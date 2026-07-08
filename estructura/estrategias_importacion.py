from abc import ABC, abstractmethod
import csv
import json
import os

class EstrategiaImportacion(ABC):
    @abstractmethod
    def importar(self, archivo, sistema):
        pass

class ImportacionCSV(EstrategiaImportacion):
    def importar(self, archivo, sistema):
        resultado = {"exitosos": 0, "duplicados": 0, "errores": []}
        
        try:
            if not os.path.exists(archivo):
                resultado["errores"].append(f"Archivo no encontrado: {archivo}")
                return resultado
            
            with open(archivo, 'r', encoding='utf-8') as f:
                muestra = f.read(1024)
                f.seek(0)
                
                if ';' in muestra:
                    delim = ';'
                elif '\t' in muestra:
                    delim = '\t'
                else:
                    delim = ','
                
                reader = csv.DictReader(f, delimiter=delim)
                
                if 'cedula' not in reader.fieldnames or 'tipo' not in reader.fieldnames:
                    resultado["errores"].append("El CSV debe tener columnas 'cedula' y 'tipo'")
                    return resultado
                
                for fila in reader:
                    try:
                        cedula = fila.get('cedula', '').strip()
                        tipo = fila.get('tipo', '').strip().lower()
                        
                        if not cedula or not tipo:
                            resultado["errores"].append(f"Datos incompletos: {fila}")
                            continue
                        
                        if tipo not in ["estudiante", "docente", "personal"]:
                            resultado["errores"].append(f"Tipo invalido '{tipo}' para cedula {cedula}")
                            continue
                        
                        if sistema.agregar_cedula_permitida(cedula, tipo):
                            resultado["exitosos"] += 1
                        else:
                            resultado["duplicados"] += 1
                    
                    except Exception as e:
                        resultado["errores"].append(f"Error procesando {fila}: {str(e)}")
        
        except Exception as e:
            resultado["errores"].append(f"Error al leer CSV: {str(e)}")
        
        return resultado

class ImportacionJSON(EstrategiaImportacion):
    def importar(self, archivo, sistema):
        resultado = {"exitosos": 0, "duplicados": 0, "errores": []}
        
        try:
            if not os.path.exists(archivo):
                resultado["errores"].append(f"Archivo no encontrado: {archivo}")
                return resultado
            
            with open(archivo, 'r', encoding='utf-8') as f:
                datos = json.load(f)
            
            if isinstance(datos, dict) and 'cedulas' in datos:
                datos = datos['cedulas']
            
            if not isinstance(datos, list):
                resultado["errores"].append("El JSON debe ser una lista de objetos o tener clave 'cedulas'")
                return resultado
            
            for item in datos:
                try:
                    if isinstance(item, dict):
                        cedula = item.get('cedula', '').strip()
                        tipo = item.get('tipo', '').strip().lower()
                    elif isinstance(item, str):
                        cedula = item.strip()
                        tipo = "estudiante"
                    else:
                        resultado["errores"].append(f"Formato invalido: {item}")
                        continue
                    
                    if not cedula:
                        resultado["errores"].append(f"Datos incompletos: {item}")
                        continue
                    
                    if tipo not in ["estudiante", "docente", "personal"]:
                        resultado["errores"].append(f"Tipo invalido '{tipo}' para cedula {cedula}")
                        continue
                    
                    if sistema.agregar_cedula_permitida(cedula, tipo):
                        resultado["exitosos"] += 1
                    else:
                        resultado["duplicados"] += 1
                
                except Exception as e:
                    resultado["errores"].append(f"Error procesando {item}: {str(e)}")
        
        except json.JSONDecodeError as e:
            resultado["errores"].append(f"Error al decodificar JSON: {str(e)}")
        except Exception as e:
            resultado["errores"].append(f"Error al leer JSON: {str(e)}")
        
        return resultado

class ImportacionExcel(EstrategiaImportacion):
    def importar(self, archivo, sistema):
        resultado = {"exitosos": 0, "duplicados": 0, "errores": []}
        
        try:
            from openpyxl import load_workbook
        except ImportError:
            resultado["errores"].append("Instalar openpyxl: pip install openpyxl")
            return resultado
        
        try:
            if not os.path.exists(archivo):
                resultado["errores"].append(f"Archivo no encontrado: {archivo}")
                return resultado
            
            wb = load_workbook(archivo, data_only=True)
            ws = wb.active
            
            if ws.max_row < 2:
                resultado["errores"].append("El Excel esta vacio o solo tiene encabezados")
                return resultado
            
            headers = []
            for cell in ws[1]:
                if cell.value:
                    headers.append(str(cell.value).strip().lower())
                else:
                    headers.append("")
            
            try:
                col_cedula = headers.index('cedula')
                col_tipo = headers.index('tipo')
            except ValueError:
                resultado["errores"].append("El Excel debe tener columnas 'cedula' y 'tipo'")
                return resultado
            
            for row in ws.iter_rows(min_row=2, values_only=True):
                try:
                    if not row or not any(row):
                        continue
                    
                    cedula = str(row[col_cedula]).strip() if col_cedula < len(row) and row[col_cedula] else ''
                    tipo = str(row[col_tipo]).strip().lower() if col_tipo < len(row) and row[col_tipo] else ''
                    
                    if not cedula or not tipo:
                        continue
                    
                    if tipo not in ["estudiante", "docente", "personal"]:
                        continue
                    
                    if sistema.agregar_cedula_permitida(cedula, tipo):
                        resultado["exitosos"] += 1
                    else:
                        resultado["duplicados"] += 1
                
                except Exception as e:
                    resultado["errores"].append(f"Error en fila {row}: {str(e)}")
        
        except Exception as e:
            resultado["errores"].append(f"Error al leer Excel: {str(e)}")
        
        return resultado

class ImportacionTXT(EstrategiaImportacion):
    def importar(self, archivo, sistema):
        resultado = {"exitosos": 0, "duplicados": 0, "errores": []}
        
        try:
            if not os.path.exists(archivo):
                resultado["errores"].append(f"Archivo no encontrado: {archivo}")
                return resultado
            
            with open(archivo, 'r', encoding='utf-8') as f:
                lineas = f.readlines()
            
            for i, linea in enumerate(lineas, 1):
                linea = linea.strip()
                if not linea or linea.startswith('#'):
                    continue
                
                partes = linea.split()
                if len(partes) >= 2:
                    cedula = partes[0].strip()
                    tipo = partes[1].strip().lower()
                elif len(partes) == 1:
                    cedula = partes[0].strip()
                    tipo = "estudiante"
                else:
                    resultado["errores"].append(f"Linea {i}: Formato invalido")
                    continue
                
                if not cedula:
                    resultado["errores"].append(f"Linea {i}: Cedula vacia")
                    continue
                
                if tipo not in ["estudiante", "docente", "personal"]:
                    resultado["errores"].append(f"Linea {i}: Tipo invalido '{tipo}'")
                    continue
                
                if sistema.agregar_cedula_permitida(cedula, tipo):
                    resultado["exitosos"] += 1
                else:
                    resultado["duplicados"] += 1
        
        except Exception as e:
            resultado["errores"].append(f"Error al leer TXT: {str(e)}")
        
        return resultado

class ImportacionAutomatica(EstrategiaImportacion):
    def importar(self, archivo, sistema):
        extension = os.path.splitext(archivo)[1].lower()
        
        if extension == '.csv':
            estrategia = ImportacionCSV()
        elif extension == '.json':
            estrategia = ImportacionJSON()
        elif extension in ['.xlsx', '.xls']:
            estrategia = ImportacionExcel()
        elif extension == '.txt':
            estrategia = ImportacionTXT()
        else:
            try:
                with open(archivo, 'r', encoding='utf-8') as f:
                    contenido = f.read(1024)
                
                if contenido.startswith('{') or contenido.startswith('['):
                    estrategia = ImportacionJSON()
                elif ',' in contenido or ';' in contenido or '\t' in contenido:
                    estrategia = ImportacionCSV()
                else:
                    estrategia = ImportacionTXT()
            except:
                return {"exitosos": 0, "duplicados": 0, "errores": ["Formato no detectable"]}
        
        return estrategia.importar(archivo, sistema)

class FabricaEstrategias:
    @staticmethod
    def crear_estrategia(formato):
        formatos = {
            'csv': ImportacionCSV,
            'json': ImportacionJSON,
            'excel': ImportacionExcel,
            'xlsx': ImportacionExcel,
            'txt': ImportacionTXT,
            'auto': ImportacionAutomatica
        }
        
        clase = formatos.get(formato.lower(), ImportacionAutomatica)
        return clase()