# Aqui se hizo Patron Strategy
# Cada clase sabe importar un tipo de archivo diferente.
# estructura/estrategias_importacion.py
"""
Importación masiva de usuarios.

Formato aceptado para CSV o Excel:
cedula,nombres,apellidos,correo,contrasena,tipo_usuario

La columna tipo_usuario puede ser: Estudiante, Docente o Personal.
El sistema crea el usuario igual que si el administrador lo registrara manualmente.
"""

from abc import ABC, abstractmethod
import csv
import os


class EstrategiaImportacion(ABC):
    @abstractmethod
    def importar(self, ruta_archivo):
        pass


def _normalizar_cedula(valor):
    return str(valor or "").strip().replace(" ", "")


def _normalizar_texto(valor):
    return str(valor or "").strip()


def _normalizar_rol(valor):
    rol = _normalizar_texto(valor).lower()
    mapa = {
        "estudiante": "Estudiante",
        "docente": "Docente",
        "personal": "Personal",
        "admin": "Personal",
        "administrador": "Personal",
    }
    return mapa.get(rol, "")


def _validar_nombres(texto):
    limpio = texto.replace(" ", "")
    return bool(limpio) and limpio.isalpha()


def _validar_encabezados(encabezados):
    requeridos = {"cedula", "nombres", "apellidos", "correo", "tipo_usuario"}
    return requeridos.issubset(set(encabezados))


def _guardar_usuario_desde_fila(fila_norm, resultado):
    # Aqui se hizo Factory indirectamente: segun el tipo_usuario se crea el rol correcto.
    from estructura.AlmacenamietoUsuarios import Almacenamiento_Usuarios
    from estructura.fabrica_usuarios import FabricaUsuarios

    cedula = _normalizar_cedula(fila_norm.get("cedula"))
    nombres = _normalizar_texto(fila_norm.get("nombres") or fila_norm.get("nombre"))
    apellidos = _normalizar_texto(fila_norm.get("apellidos") or fila_norm.get("apellido"))
    correo = _normalizar_texto(fila_norm.get("correo")).lower()
    contrasena = _normalizar_texto(fila_norm.get("contrasena") or fila_norm.get("contraseña") or "1234")
    rol = _normalizar_rol(fila_norm.get("tipo_usuario") or fila_norm.get("rol"))

    if not (cedula and nombres and apellidos and correo and contrasena and rol):
        return False, "datos incompletos o tipo_usuario inválido"
    if not (cedula.isdigit() and len(cedula) == 10):
        return False, "la cédula debe tener exactamente 10 dígitos"
    if "@" not in correo or correo.endswith("@") or correo.startswith("@"):
        return False, "correo inválido"
    if not _validar_nombres(nombres):
        return False, "los nombres solo deben contener letras y espacios"
    if not _validar_nombres(apellidos):
        return False, "los apellidos solo deben contener letras y espacios"

    almacenamiento = Almacenamiento_Usuarios()
    usuarios = almacenamiento.json.repo.leer_todo()

    if any(u.get("cedula") == cedula for u in usuarios):
        resultado["duplicados"] += 1
        return True, "duplicado"
    if any(str(u.get("correo", "")).lower() == correo for u in usuarios):
        resultado["duplicados"] += 1
        return True, "duplicado"

    datos_usuario = {
        "cedula": cedula,
        "nombre": nombres,
        "apellido": apellidos,
        "correo": correo,
        "contraseña": contrasena,
        "rol": rol,
    }

    # Aqui se hizo Patron Factory
    usuario_objeto = FabricaUsuarios().crear_usuario(datos_usuario)
    if almacenamiento.guardar_usuario_objeto(usuario_objeto):
        return True, "usuario creado"
    resultado["duplicados"] += 1
    return True, "duplicado"


class ImportacionCSV(EstrategiaImportacion):
    def importar(self, ruta_archivo):
        resultado = {"total": 0, "exitosos": 0, "duplicados": 0, "errores": []}

        if not os.path.exists(ruta_archivo):
            resultado["errores"].append(f"Archivo no encontrado: {ruta_archivo}")
            return resultado

        with open(ruta_archivo, "r", encoding="utf-8-sig") as f:
            muestra = f.read(1024)
            f.seek(0)
            delimitador = ";" if ";" in muestra else ("\t" if "\t" in muestra else ",")
            lector = csv.DictReader(f, delimiter=delimitador)
            campos = [(c or "").strip().lower() for c in (lector.fieldnames or [])]

            if not _validar_encabezados(campos):
                resultado["errores"].append(
                    "El archivo debe tener estas columnas: cedula,nombres,apellidos,correo,contrasena,tipo_usuario"
                )
                return resultado

            for num_fila, fila in enumerate(lector, start=2):
                resultado["total"] += 1
                fila_norm = {(k or "").strip().lower(): (v or "").strip() for k, v in fila.items()}
                ok, msg = _guardar_usuario_desde_fila(fila_norm, resultado)
                if ok and msg == "usuario creado":
                    resultado["exitosos"] += 1
                elif not ok:
                    resultado["errores"].append(f"Fila {num_fila}: {msg}")

        return resultado


class ImportacionExcel(EstrategiaImportacion):
    def importar(self, ruta_archivo):
        resultado = {"total": 0, "exitosos": 0, "duplicados": 0, "errores": []}

        try:
            from openpyxl import load_workbook
        except ImportError:
            resultado["errores"].append("Falta instalar 'openpyxl' para importar Excel.")
            return resultado

        if not os.path.exists(ruta_archivo):
            resultado["errores"].append(f"Archivo no encontrado: {ruta_archivo}")
            return resultado

        wb = load_workbook(ruta_archivo, data_only=True)
        ws = wb.active
        encabezados = [str(celda.value).strip().lower() if celda.value else "" for celda in ws[1]]

        if not _validar_encabezados(encabezados):
            resultado["errores"].append(
                "El Excel debe tener estas columnas: cedula,nombres,apellidos,correo,contrasena,tipo_usuario"
            )
            return resultado

        for idx, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not fila or not any(fila):
                continue
            resultado["total"] += 1
            fila_norm = {
                encabezados[i]: str(fila[i]).strip() if i < len(fila) and fila[i] is not None else ""
                for i in range(len(encabezados))
            }
            ok, msg = _guardar_usuario_desde_fila(fila_norm, resultado)
            if ok and msg == "usuario creado":
                resultado["exitosos"] += 1
            elif not ok:
                resultado["errores"].append(f"Fila {idx}: {msg}")

        return resultado


# Aqui se hizo Patron Factory
class FabricaEstrategiasImportacion:
    @staticmethod
    def crear_para_archivo(ruta_archivo):
        extension = os.path.splitext(ruta_archivo)[1].lower()
        if extension == ".csv":
            return ImportacionCSV()
        if extension in (".xlsx", ".xls"):
            return ImportacionExcel()
        raise ValueError(f"Formato de archivo no soportado: '{extension}'")
